import csv
import copy
import ipaddress
import json
import math
import queue
import random
import socket
import struct
import threading
import time
import webbrowser
from collections import deque
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path
import tkinter as tk
from tkinter import ttk, messagebox

import numpy as np
from openpyxl import load_workbook
from scipy import signal
import serial

import matplotlib
matplotlib.use("TkAgg")
from matplotlib.backends.backend_tkagg import FigureCanvasTkAgg
from matplotlib.figure import Figure

BASE = Path(__file__).resolve().parent

@dataclass
class WaveResult:
    hs_m: float = float("nan")
    tp_s: float = float("nan")
    valid: bool = False
    timestamp: float = 0.0

# ==========================================
# GESTION DES DONNÉES (CSV & RAO)
# ==========================================

class CSVLogger:
    """Write measurements and errors to a comma-separated log."""
    def __init__(self, directory):
        self.directory = Path(directory)
        self.directory.mkdir(parents=True, exist_ok=True)
        stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        self.log_path = self.directory / f"asn_measurements_{stamp}.csv"
        self.file = None
        self.writer = None

    def start(self):
        self.file = self.log_path.open("w", encoding="utf-8", newline="")
        self.writer = csv.writer(self.file, lineterminator="\n")
        self.writer.writerow((
            "timestamp_utc", "type", "heave_m", "current_speed_kn",
            "current_direction_deg", "current_depth_m", "wind_speed_kn",
            "wind_direction_deg", "hs_m", "tp_s", "level", "component", "message",
        ))
        self.file.flush()

    def write(self, data):
        if self.writer and self.file:
            self.writer.writerow((data[0], "measurement", *data[1:], "", "", ""))
            self.file.flush()

    def write_error(self, level, component, message):
        if self.writer and self.file:
            timestamp = datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")
            self.writer.writerow((timestamp, "error", *([""] * 8), level, component, message))
            self.file.flush()

    def close(self):
        if self.file:
            self.file.close()
            self.file = None
            self.writer = None


def build_input_configuration(current_inputs, acquisition_mode, equipment_settings):
    """Validate UI port settings and preserve the active port expected by receivers."""
    updated_inputs = copy.deepcopy(current_inputs)
    if acquisition_mode not in ("Capteurs", "Simulation"):
        raise ValueError("Mode d'acquisition invalide.")
    updated_inputs["mode"] = "simulation" if acquisition_mode == "Simulation" else "udp"

    for equipment, settings in equipment_settings.items():
        if equipment not in updated_inputs:
            continue
        config = updated_inputs[equipment]
        previous_transport = config.get("transport", "udp")
        transport = settings["transport"]
        if transport not in ("udp", "serial"):
            raise ValueError(f"Transport invalide pour {equipment}.")

        source_host = settings.get("udp_host", "").strip()
        if transport == "udp" and source_host:
            try:
                source_host = str(ipaddress.IPv4Address(source_host))
            except ipaddress.AddressValueError as exc:
                raise ValueError(f"L'adresse IP source de {equipment} doit être une adresse IPv4 valide.") from exc

        udp_port_text = settings["udp_port"].strip()
        if udp_port_text:
            try:
                udp_port = int(udp_port_text)
            except ValueError as exc:
                raise ValueError(f"Le port UDP de {equipment} doit être un nombre.") from exc
            if not 1 <= udp_port <= 65535:
                raise ValueError(f"Le port UDP de {equipment} doit être compris entre 1 et 65535.")
        else:
            udp_port = config.get("udp_port")
            if udp_port is None and previous_transport == "udp":
                udp_port = config.get("port")
        if transport == "udp" and udp_port is None:
            raise ValueError(f"Renseignez le port UDP de {equipment}.")

        serial_port = settings["serial_port"].strip()
        if not serial_port:
            serial_port = config.get("serial_port", "")
            if not serial_port and previous_transport == "serial":
                serial_port = str(config.get("port", ""))
        if transport == "serial" and not serial_port:
            raise ValueError(f"Renseignez le port série de {equipment} (par exemple COM3).")

        baudrate_text = settings["baudrate"].strip()
        try:
            baudrate = int(baudrate_text or config.get("baudrate", 115200))
        except ValueError as exc:
            raise ValueError(f"Le débit série de {equipment} doit être un nombre.") from exc
        if baudrate <= 0:
            raise ValueError(f"Le débit série de {equipment} doit être supérieur à zéro.")

        config["transport"] = transport
        config["host"] = source_host
        config["udp_port"] = udp_port if udp_port is not None else ""
        config["serial_port"] = serial_port
        config["baudrate"] = baudrate
        config["port"] = udp_port if transport == "udp" else serial_port

    return updated_inputs


def determine_reception_quality(valid, missing_sensors=0, error_count=0):
    """Return a coarse quality label based on signal validity and sensor health."""
    if not valid and (missing_sensors >= 2 or error_count >= 3):
        return "POOR"
    if missing_sensors > 0 or error_count > 0:
        return "MODERATE"
    return "GOOD"

class RAOManager:
    """Gestionnaire de RAO (Response Amplitude Operator)."""
    def __init__(self, filepath=None, heading_deg=0):
        self.curves = {}
        self.heading_deg = float(heading_deg)
        self.load_error = None
        if filepath:
            try:
                self._load_rao(filepath)
            except (OSError, ValueError, KeyError, TypeError) as exc:
                self.load_error = str(exc)

    @property
    def available_headings(self):
        return tuple(sorted(self.curves))

    def _load_rao(self, filepath):
        path = Path(filepath)
        if not path.is_file():
            raise FileNotFoundError(f"Fichier RAO introuvable: {path}")

        if path.suffix.lower() == ".xlsx":
            workbook = load_workbook(path, read_only=True, data_only=True)
            try:
                for sheet in workbook.worksheets:
                    sheet_name = sheet.title.strip()
                    if sheet_name.endswith("°"):
                        sheet_name = sheet_name[:-1].strip()
                    try:
                        heading = float(sheet_name)
                    except ValueError:
                        continue

                    points = []
                    if sheet["A4"].value == "Period" and sheet["F4"].value == "Heave":
                        rows = sheet.iter_rows(min_row=7, min_col=1, max_col=6, values_only=True)
                        period_index, gain_index = 0, 5
                    else:
                        rows = sheet.iter_rows(min_row=5, min_col=2, max_col=5, values_only=True)
                        period_index, gain_index = 0, 3
                    for row in rows:
                        period, gain = row[period_index], row[gain_index]
                        if period is None or gain is None:
                            continue
                        period, gain = float(period), float(gain)
                        if math.isfinite(period) and period > 0 and math.isfinite(gain) and gain >= 0:
                            points.append((1.0 / period, gain))
                    if points:
                        points.sort(key=lambda point: point[0])
                        self.curves[heading] = (
                            np.asarray([point[0] for point in points]),
                            np.asarray([point[1] for point in points]),
                        )
            finally:
                workbook.close()
        elif path.suffix.lower() == ".csv":
            points = []
            with path.open("r", encoding="utf-8-sig", newline="") as source:
                for row in csv.reader(source):
                    if len(row) < 2:
                        continue
                    try:
                        frequency = float(row[0].strip().replace(",", "."))
                        gain = float(row[1].strip().replace(",", "."))
                    except ValueError:
                        continue
                    if math.isfinite(frequency) and frequency > 0 and math.isfinite(gain) and gain >= 0:
                        points.append((frequency, gain))
            if points:
                points.sort(key=lambda point: point[0])
                self.curves[0.0] = (
                    np.asarray([point[0] for point in points]),
                    np.asarray([point[1] for point in points]),
                )
        else:
            raise ValueError(f"Format RAO non pris en charge: {path.suffix}")

        if not self.curves:
            raise ValueError(f"Aucune courbe RAO exploitable dans {path.name}")

    def set_heading(self, heading_deg):
        self.heading_deg = float(heading_deg)

    def get_gain(self, frequency):
        if not self.curves:
            return 1.0
        heading = min(self.curves, key=lambda value: abs(value - self.heading_deg))
        frequencies, gains = self.curves[heading]
        if frequency < frequencies[0] or frequency > frequencies[-1]:
            return 1.0
        return float(np.interp(frequency, frequencies, gains))

# ==========================================
# ACQUISITION (Réseau, Série, Parseurs)
# ==========================================

class NMEAParser:
    @staticmethod
    def text(payload):
        return payload.decode("ascii", errors="ignore").strip()

    @staticmethod
    def heave(payload):
        s = NMEAParser.text(payload)
        parts = s.replace("*", ",").split(",")
        if parts and parts[0].upper().endswith("HEAVE") and len(parts) >= 2:
            return ("heave", float(parts[1]))
        return None
    # (Les autres parseurs NMEA restent similaires...)
    @staticmethod
    def current(payload):
        s = NMEAParser.text(payload)
        parts = s.replace("*", ",").split(",")
        if parts and parts[0].upper() in ("CURRENT", "CUR") and len(parts) >= 4:
            return ("current", float(parts[1]), float(parts[2]), float(parts[3]))
        return None

    @staticmethod
    def wind(payload):
        s = NMEAParser.text(payload)
        parts = s.replace("*", ",").split(",")
        if parts and parts[0].upper() in ("WIND", "MWV") and len(parts) >= 3:
            return ("wind", float(parts[1]), float(parts[2]))
        return None

class ExailParser:
    """Decodeur minimal du format binaire STDBIN Exail."""
    @staticmethod
    def phlin(payload):
        """Decode a PHLIN sentence from a single or multi-sentence payload."""
        try:
            sentences = payload.decode("ascii").splitlines()
        except UnicodeDecodeError:
            return None

        for sentence in sentences:
            sentence = sentence.strip()
            if not sentence.upper().startswith("$PHLIN,") or "*" not in sentence:
                continue
            try:
                body, checksum = sentence[1:].rsplit("*", 1)
                if len(checksum) != 2:
                    continue
                expected = 0
                for character in body.encode("ascii"):
                    expected ^= character
                if int(checksum, 16) != expected:
                    continue

                fields = body.split(",")
                if len(fields) != 4 or fields[0].upper() != "PHLIN":
                    continue
                return ("heave", float(fields[3]))
            except ValueError:
                continue
        return None

    @staticmethod
    def stdbin(payload):
        """Decode une trame STDBIN complete et retourne le heave en metres."""
        if len(payload) < 32 or payload[:2] != b"IX":
            return None

        try:
            offset = 2
            version = payload[offset]
            offset += 1
            nav_mask = int.from_bytes(payload[offset:offset + 4], "big")
            offset += 4
            offset += 8  # extended navigation mask and extended mask
            if version > 3:
                offset += 2  # navigation block size
            total_size = int.from_bytes(payload[offset:offset + 2], "big")
            offset += 2

            if total_size < offset + 8 or len(payload) < total_size:
                return None

            offset += 8  # INS timestamp and telegram counter
            if nav_mask & (1 << 0):
                offset += 12
            if nav_mask & (1 << 1):
                offset += 12
            if not nav_mask & (1 << 2) or offset + 16 > total_size:
                return None

            heave = struct.unpack_from(">4f", payload, offset)[1]
            return ("heave", float(heave))
        except (IndexError, struct.error, ValueError):
            return None

    @staticmethod
    def bacustom2(payload):
        """Alias retained for configurations naming the Exail output BACUSTOM2."""
        return ExailParser.stdbin(payload)


class StdBinStreamParser:
    """Accumulate un flux binaire et extrait les trames STDBIN completes."""
    def __init__(self):
        self.buffer = bytearray()

    def feed(self, payload):
        self.buffer.extend(payload)
        results = []
        while True:
            sync_idx = self.buffer.find(b"IX")
            if sync_idx < 0:
                if self.buffer.endswith(b"I"):
                    self.buffer[:] = b"I"
                else:
                    self.buffer.clear()
                break
            if sync_idx:
                del self.buffer[:sync_idx]
            if len(self.buffer) < 17:
                break

            version = self.buffer[2]
            size_offset = 15 if version <= 3 else 17
            if len(self.buffer) < size_offset + 2:
                break
            total_size = int.from_bytes(self.buffer[size_offset:size_offset + 2], "big")
            if total_size < size_offset + 2:
                del self.buffer[:2]
                continue
            if len(self.buffer) < total_size:
                break

            frame = bytes(self.buffer[:total_size])
            del self.buffer[:total_size]
            item = ExailParser.stdbin(frame)
            if item:
                results.append(item)
        return results


class PhlinStreamParser:
    """Accumulate un flux texte et extrait les trames $PHLIN completes."""
    def __init__(self):
        self.buffer = bytearray()

    def feed(self, payload):
        self.buffer.extend(payload)
        results = []
        while b"\n" in self.buffer:
            line, _, remainder = self.buffer.partition(b"\n")
            self.buffer = bytearray(remainder)
            item = ExailParser.phlin(line)
            if item:
                results.append(item)
        return results

class UDPReceiver(threading.Thread):
    def __init__(self, port, out_queue, decoder, bind_host="0.0.0.0", source_host=None):
        super().__init__(daemon=True)
        self.port = port
        self.out_queue = out_queue
        self.decoder = decoder
        self.bind_host = bind_host
        self.source_host = source_host
        self.stop_event = threading.Event()

    def run(self):
        sock = socket.socket(socket.AF_INET, socket.SOCK_DGRAM)
        sock.setsockopt(socket.SOL_SOCKET, socket.SO_REUSEADDR, 1)
        sock.bind((self.bind_host, self.port))
        sock.settimeout(0.5)
        try:
            while not self.stop_event.is_set():
                try:
                    payload, source = sock.recvfrom(65535)
                    if self.source_host and source[0] != self.source_host:
                        continue
                    item = self.decoder(payload)
                    if item:
                        self.out_queue.put((*item, time.time()))
                except socket.timeout:
                    continue
                except Exception as exc:
                    self.out_queue.put(("error", f"UDP {self.port}: {exc}"))
        finally:
            sock.close()
            
    def stop(self):
        self.stop_event.set()

class SerialReceiver(threading.Thread):
    """Acquisition sur port série (RS232/RS422)."""
    def __init__(self, port, baudrate, out_queue, decoder, stream_parser=None):
        super().__init__(daemon=True)
        self.port = port
        self.baudrate = baudrate
        self.out_queue = out_queue
        self.decoder = decoder
        self.stream_parser = stream_parser
        self.stop_event = threading.Event()

    def run(self):
        try:
            with serial.Serial(self.port, self.baudrate, timeout=0.5) as ser:
                while not self.stop_event.is_set():
                    payload = ser.read(ser.in_waiting or 1)
                    if not payload:
                        continue
                    if self.stream_parser:
                        items = self.stream_parser.feed(payload)
                    else:
                        item = self.decoder(payload)
                        items = [item] if item else []
                    for item in items:
                        self.out_queue.put((*item, time.time()))
        except Exception as exc:
            self.out_queue.put(("error", f"Serial {self.port}: {exc}"))

    def stop(self):
        self.stop_event.set()

# ==========================================
# TRAITEMENT DU SIGNAL EN ARRIÈRE-PLAN
# ==========================================

def compute_wave(
    heaves,
    sample_rate_hz,
    rao_manager=None,
    fmin=0.03,
    fmax=0.5,
    min_duration_seconds=60,
):
    """Fonction mathématique pure, isolée des threads."""
    x = np.asarray(heaves, dtype=float)
    if len(x) < max(32, int(sample_rate_hz * min_duration_seconds)):
        return WaveResult()

    x = signal.detrend(x)
    nperseg = min(len(x), max(256, int(sample_rate_hz * 256)))
    freqs, psd = signal.welch(x, fs=sample_rate_hz, window="hann", nperseg=nperseg, detrend="constant")

    mask = (freqs >= fmin) & (freqs <= fmax)
    freqs, psd = freqs[mask], psd[mask]

    if len(freqs) < 3:
        return WaveResult()

    if rao_manager is None:
        gain = np.ones_like(freqs, dtype=float)
    else:
        gain = np.asarray([max(float(rao_manager.get_gain(f)), 1e-6) for f in freqs])
    psd = psd / (gain ** 2)

    m0 = np.trapezoid(psd, freqs) if hasattr(np, "trapezoid") else np.trapz(psd, freqs)
    if not np.isfinite(m0) or m0 <= 0:
        return WaveResult()

    peak_idx = int(np.argmax(psd))
    fp = float(freqs[peak_idx])
    tp = 1.0 / fp if fp > 0 else float("nan")
    hs = 4.0 * math.sqrt(m0)
    return WaveResult(hs, tp, True, time.time())

class WaveProcessor(threading.Thread):
    """Thread dédié pour ne pas bloquer l'interface UI (Tkinter)."""
    def __init__(self, data_lock, heave_deque, sample_rate, result_queue, config, rao_manager):
        super().__init__(daemon=True)
        self.data_lock = data_lock
        self.heave_deque = heave_deque
        self.sample_rate = sample_rate
        self.result_queue = result_queue
        self.config = config
        self.rao_manager = rao_manager
        self.stop_event = threading.Event()

    def run(self):
        while not self.stop_event.is_set():
            # Copie rapide sous verrou pour minimiser le blocage
            with self.data_lock:
                heave_copy = list(self.heave_deque)
            
            if heave_copy:
                rao_manager = (
                    None if self.config["inputs"].get("mode") == "simulation" else self.rao_manager
                )
                result = compute_wave(
                    heave_copy, 
                    self.sample_rate, 
                    rao_manager,
                    fmin=self.config["wave"]["min_frequency_hz"],
                    fmax=self.config["wave"]["max_frequency_hz"],
                    min_duration_seconds=(
                        10 if self.config["inputs"].get("mode") == "simulation" else 60
                    ),
                )
                self.result_queue.put(result)
            
            # Ne recalcule le spectre que toutes les 2 secondes
            time.sleep(2.0)

    def stop(self):
        self.stop_event.set()

# ==========================================
# SIMULATEUR & INTERFACE GRAPHIQUE
# ==========================================

SEA_STATE_PRESETS = {
    "calm": {
        "swell_h": 0.18,
        "chop_h": 0.08,
        "drift_h": 0.05,
        "noise_std": 0.015,
        "current_speed": 0.5,
        "wind_speed": 6.0,
    },
    "moderate": {
        "swell_h": 0.42,
        "chop_h": 0.15,
        "drift_h": 0.10,
        "noise_std": 0.025,
        "current_speed": 1.1,
        "wind_speed": 14.0,
    },
    "rough": {
        "swell_h": 0.85,
        "chop_h": 0.30,
        "drift_h": 0.18,
        "noise_std": 0.04,
        "current_speed": 2.0,
        "wind_speed": 25.0,
    },
    "instrument": {
        "swell_h": 0.25,
        "chop_h": 0.12,
        "drift_h": 0.05,
        "noise_std": 0.08,
        "current_speed": 1.0,
        "wind_speed": 12.0,
    },
}


def simulate_sea_state(time_s, phase=0.0, preset="moderate"):
    """Generate a configurable sea-state signal with swell, chop, drift and noise."""
    config = SEA_STATE_PRESETS.get(preset, SEA_STATE_PRESETS["moderate"])
    swell = config["swell_h"] * math.sin(2 * math.pi * 0.09 * time_s + phase)
    chop = config["chop_h"] * math.sin(2 * math.pi * 0.18 * time_s + phase * 1.7)
    low_freq_drift = config["drift_h"] * math.sin(2 * math.pi * 0.03 * time_s + phase * 0.6)
    modulation = 1.0 + 0.25 * math.sin(2 * math.pi * 0.04 * time_s + phase * 0.8)
    noise = random.gauss(0.0, config["noise_std"])
    return (swell + chop + low_freq_drift) * modulation + noise


class Simulator(threading.Thread):
    """Générateur de données plus réaliste pour le mode simulation."""
    def __init__(self, out_queue, sample_rate=5.0, preset="moderate"):
        super().__init__(daemon=True)
        self.out_queue = out_queue
        self.sample_rate = sample_rate
        self.preset = preset
        self.stop_event = threading.Event()
        self.t = 0.0
        self.phase = random.random() * 6.28
        self.config = SEA_STATE_PRESETS.get(preset, SEA_STATE_PRESETS["moderate"])

    def run(self):
        dt = 1.0 / self.sample_rate
        while not self.stop_event.is_set():
            wave = simulate_sea_state(self.t, phase=self.phase, preset=self.preset)
            current_speed = self.config["current_speed"] * (1.0 + 0.20 * math.sin(2 * math.pi * 0.02 * self.t))
            wind_speed = self.config["wind_speed"] * (1.0 + 0.18 * math.sin(2 * math.pi * 0.015 * self.t))

            now = time.time()
            self.out_queue.put(("heave", wave, now))
            self.out_queue.put(("current", current_speed, 210, 2.5, now))
            self.out_queue.put(("wind", wind_speed, 250, now))

            self.t += dt
            time.sleep(dt)

    def stop(self):
        self.stop_event.set()

class App(tk.Tk):
    def __init__(self):
        super().__init__()
        self.title("ASN Wave Monitor")
        self.geometry("1250x780")
        self.config_data = json.loads((BASE / "config.json").read_text(encoding="utf-8"))
        
        self.queue = queue.Queue()
        self.result_queue = queue.Queue()
        self.data_lock = threading.Lock()
        
        self.sample_rate = 5.0
        self.heave = deque(maxlen=int(self.sample_rate * 20 * 60))
        
        self.history_t = deque(maxlen=1800)
        self.history_hs = deque(maxlen=1800)
        self.history_tp = deque(maxlen=1800)
        self.history_current = deque(maxlen=1800)
        self.history_wind = deque(maxlen=1800)
        
        self.current = {"speed": float("nan"), "direction": float("nan"), "depth": float("nan")}
        self.wind = {"speed": float("nan"), "direction": float("nan")}
        self.wave = WaveResult()

        self.vessels = self.config_data.get("vessels") or {
            "Profil actuel": {"rao_file": self.config_data["wave"].get("rao_file")}
        }
        self.selected_vessel = self.config_data.get("selected_vessel")
        if self.selected_vessel not in self.vessels:
            self.selected_vessel = next(iter(self.vessels))
        self.rao = self._load_vessel_rao(self.selected_vessel)

        self.data_directory = BASE / self.config_data["storage"]["directory"]
        self.logger = CSVLogger(self.data_directory)
        self.logger.start()

        self.build_ui()
        self.start_inputs()
        self.after(self.config_data["app"]["refresh_ms"], self.tick)
        self.protocol("WM_DELETE_WINDOW", self.on_close)

    def build_ui(self):
        """Construction de l'interface (simplifiée ici pour la lisibilité)."""
        top = ttk.Frame(self, padding=10)
        top.pack(fill="x")
        ttk.Label(top, text="ASN Wave Monitor", font=("Segoe UI", 20, "bold")).pack(side="left")
        self.status_var = tk.StringVar(value="Démarrage…")
        ttk.Label(top, textvariable=self.status_var).pack(side="right")

        self.notebook = ttk.Notebook(self)
        self.notebook.pack(fill="both", expand=True)
        dashboard = ttk.Frame(self.notebook)
        data_tab = ttk.Frame(self.notebook)
        graphs_tab = ttk.Frame(self.notebook)
        ports_tab = ttk.Frame(self.notebook)
        self.notebook.add(dashboard, text="Tableau de bord")
        self.notebook.add(data_tab, text="Données")
        self.notebook.add(graphs_tab, text="Graphiques")
        self.notebook.add(ports_tab, text="Ports")
        self.notebook.bind("<<NotebookTabChanged>>", self.on_tab_changed)

        selection = ttk.Frame(dashboard, padding=(12, 0, 12, 8))
        selection.pack(fill="x")
        ttk.Label(selection, text="Navire").pack(side="left", padx=(0, 6))
        self.vessel_var = tk.StringVar(value=self.selected_vessel)
        self.vessel_combo = ttk.Combobox(
            selection, textvariable=self.vessel_var, state="readonly",
            values=list(self.vessels), width=30
        )
        self.vessel_combo.pack(side="left", padx=(0, 16))
        self.vessel_combo.bind("<<ComboboxSelected>>", self.on_vessel_select)

        ttk.Label(selection, text="Direction de houle").pack(side="left", padx=(0, 6))
        self.heading_var = tk.StringVar()
        self.heading_combo = ttk.Combobox(
            selection, textvariable=self.heading_var, state="disabled", width=10
        )
        self.heading_combo.pack(side="left", padx=(0, 16))
        self.heading_combo.bind("<<ComboboxSelected>>", self.on_heading_select)
        self.rao_status_var = tk.StringVar()
        ttk.Label(selection, textvariable=self.rao_status_var).pack(side="left")
        self._configure_heading_control()

        cards = ttk.Frame(dashboard, padding=10)
        cards.pack(fill="x")
        self.vars = {
            "Hs": (tk.StringVar(value="—"), "m"),
            "Tp": (tk.StringVar(value="—"), "s"),
            "Courant": (tk.StringVar(value="—"), "kn"),
            "Vent": (tk.StringVar(value="—"), "kn")
        }
        for title, (var, unit) in self.vars.items():
            f = ttk.LabelFrame(cards, text=title, padding=12)
            f.pack(side="left", fill="both", expand=True, padx=5)
            ttk.Label(f, textvariable=var, font=("Segoe UI", 22, "bold")).pack()
            ttk.Label(f, text=unit).pack()

        self.alert_var = tk.StringVar(value="NORMAL")
        self.quality_var = tk.StringVar(value="GOOD")
        self.error_log_messages = deque(maxlen=20)
        self.sensor_status_vars = {}
        self.sensor_status_labels = {}

        alert = ttk.LabelFrame(dashboard, text="État des seuils", padding=8)
        alert.pack(fill="x", padx=15)
        ttk.Label(alert, textvariable=self.alert_var, font=("Segoe UI", 14, "bold")).pack(side="left")

        quality = ttk.LabelFrame(dashboard, text="Qualité de réception", padding=8)
        quality.pack(fill="x", padx=15, pady=(0, 8))
        self.quality_label = ttk.Label(quality, textvariable=self.quality_var, font=("Segoe UI", 14, "bold"))
        self.quality_label.pack(side="left")

        bottom = ttk.Frame(dashboard)
        bottom.pack(fill="both", expand=True, padx=15, pady=(0, 10))

        sensor_panel = ttk.LabelFrame(bottom, text="État des capteurs", padding=8)
        sensor_panel.pack(side="left", fill="y", padx=(0, 8))
        self.sensor_listbox = tk.Listbox(sensor_panel, width=24, height=5, exportselection=False)
        self.sensor_listbox.pack(fill="y", expand=True)
        self.sensor_listbox.insert(tk.END, "Octans")
        self.sensor_listbox.insert(tk.END, "Courant")
        self.sensor_listbox.insert(tk.END, "Vent")
        self.sensor_listbox.bind("<<ListboxSelect>>", self.on_sensor_select)

        self.sensor_details = tk.StringVar(value="Sélectionnez un capteur")
        ttk.Label(sensor_panel, textvariable=self.sensor_details, wraplength=220, justify="left").pack(fill="x", pady=(8, 0))

        log = ttk.LabelFrame(bottom, text="Journal d'erreurs", padding=8)
        log.pack(side="left", fill="both", expand=True)
        self.error_log_box = tk.Text(log, height=8, wrap="word", state="disabled", bg="#f7f7f7")
        self.error_log_box.pack(side="left", fill="both", expand=True)
        yscroll = tk.Scrollbar(log, orient="vertical", command=self.error_log_box.yview)
        yscroll.pack(side="right", fill="y")
        self.error_log_box.configure(yscrollcommand=yscroll.set)

        self._build_graphs_tab(graphs_tab)
        self._build_data_tab(data_tab)
        self.refresh_data_files()
        self._build_input_settings_tab(ports_tab)

    def _build_graphs_tab(self, parent):
        graph_frame = ttk.Frame(parent, padding=12)
        graph_frame.pack(fill="both", expand=True)
        self.fig = Figure(figsize=(13, 8), dpi=100)
        self.wave_ax = self.fig.add_subplot(211)
        self.tp_ax = self.wave_ax.twinx()
        self.environment_ax = self.fig.add_subplot(212, sharex=self.wave_ax)
        self.wind_ax = self.environment_ax.twinx()
        self.fig.subplots_adjust(left=0.08, right=0.9, top=0.95, bottom=0.1, hspace=0.45)
        self.canvas = FigureCanvasTkAgg(self.fig, master=graph_frame)
        self.canvas.get_tk_widget().pack(fill="both", expand=True)

    def _build_data_tab(self, parent):
        toolbar = ttk.Frame(parent, padding=10)
        toolbar.pack(fill="x")
        ttk.Button(toolbar, text="Actualiser", command=self.refresh_data_files).pack(side="left")
        ttk.Button(toolbar, text="Ouvrir le fichier", command=self.open_selected_data_file).pack(
            side="left", padx=(8, 12)
        )
        self.data_status_var = tk.StringVar(value="Sélectionnez un fichier de données.")
        ttk.Label(toolbar, textvariable=self.data_status_var).pack(side="left")

        panes = ttk.Panedwindow(parent, orient=tk.HORIZONTAL)
        panes.pack(fill="both", expand=True, padx=10, pady=(0, 10))

        files_panel = ttk.LabelFrame(panes, text="Fichiers enregistrés", padding=8)
        panes.add(files_panel, weight=1)
        self.data_file_tree = ttk.Treeview(
            files_panel, columns=("modified", "size"), show="tree headings", selectmode="browse"
        )
        self.data_file_tree.heading("#0", text="Fichier")
        self.data_file_tree.heading("modified", text="Modifié")
        self.data_file_tree.heading("size", text="Taille")
        self.data_file_tree.column("#0", width=220, minwidth=150)
        self.data_file_tree.column("modified", width=135, minwidth=120, stretch=False)
        self.data_file_tree.column("size", width=75, minwidth=65, stretch=False, anchor="e")
        self.data_file_tree.pack(fill="both", expand=True)
        self.data_file_tree.bind("<<TreeviewSelect>>", self.on_data_file_select)

        preview_panel = ttk.LabelFrame(panes, text="Enregistrements récents", padding=8)
        panes.add(preview_panel, weight=4)
        columns = (
            "type", "timestamp_utc", "heave_m", "current_speed_kn",
            "current_direction_deg", "current_depth_m", "wind_speed_kn",
            "wind_direction_deg", "hs_m", "tp_s", "level", "component", "message",
        )
        labels = {
            "type": "Type", "timestamp_utc": "Horodatage UTC", "heave_m": "Heave (m)",
            "current_speed_kn": "Courant (kn)", "current_direction_deg": "Dir. courant (°)",
            "current_depth_m": "Profondeur (m)", "wind_speed_kn": "Vent (kn)",
            "wind_direction_deg": "Dir. vent (°)", "hs_m": "Hs (m)", "tp_s": "Tp (s)",
            "level": "Niveau", "component": "Composant", "message": "Message",
        }
        self.data_tree = ttk.Treeview(preview_panel, columns=columns, show="headings")
        for column in columns:
            self.data_tree.heading(column, text=labels[column])
            self.data_tree.column(column, width=105, minwidth=80, stretch=False)
        self.data_tree.column("timestamp_utc", width=170, minwidth=160)
        self.data_tree.column("type", width=85, minwidth=75)
        self.data_tree.column("message", width=240, minwidth=180)
        vertical_scroll = ttk.Scrollbar(preview_panel, orient="vertical", command=self.data_tree.yview)
        horizontal_scroll = ttk.Scrollbar(preview_panel, orient="horizontal", command=self.data_tree.xview)
        self.data_tree.configure(yscrollcommand=vertical_scroll.set, xscrollcommand=horizontal_scroll.set)
        preview_panel.rowconfigure(0, weight=1)
        preview_panel.columnconfigure(0, weight=1)
        self.data_tree.grid(row=0, column=0, sticky="nsew")
        vertical_scroll.grid(row=0, column=1, sticky="ns")
        horizontal_scroll.grid(row=1, column=0, sticky="ew")

    def on_tab_changed(self, event=None):
        if self.notebook.select() == str(self.notebook.tabs()[-1]):
            self.refresh_data_files()

    def refresh_data_files(self):
        selected = self.data_file_tree.selection()
        selected_path = selected[0] if selected else None
        for item in self.data_file_tree.get_children():
            self.data_file_tree.delete(item)

        paths = sorted(
            list(self.data_directory.glob("asn_measurements_*.csv"))
            + list(self.data_directory.glob("asn_measurements_*.txt")),
            key=lambda path: path.stat().st_mtime,
            reverse=True,
        )
        for path in paths:
            modified = datetime.fromtimestamp(path.stat().st_mtime).strftime("%Y-%m-%d %H:%M:%S")
            size = f"{path.stat().st_size / 1024:.1f} KB"
            self.data_file_tree.insert(
                "", tk.END, iid=str(path.resolve()), text=path.name, values=(modified, size)
            )

        if not paths:
            self._clear_data_preview()
            self.data_status_var.set("Aucun journal trouvé dans le dossier data.")
            return

        item = selected_path if selected_path in self.data_file_tree.get_children() else str(paths[0].resolve())
        self.data_file_tree.selection_set(item)
        self.data_file_tree.focus(item)
        self.show_data_file(Path(item))

    def on_data_file_select(self, event=None):
        selection = self.data_file_tree.selection()
        if selection:
            self.show_data_file(Path(selection[0]))

    def show_data_file(self, filepath):
        self._clear_data_preview()
        recent_rows = deque(maxlen=500)
        row_count = 0
        try:
            with filepath.open("r", encoding="utf-8", newline="") as source:
                delimiter = "\t" if filepath.suffix.lower() == ".txt" else ","
                reader = csv.DictReader(source, delimiter=delimiter)
                for row in reader:
                    recent_rows.append(row)
                    row_count += 1
        except (OSError, csv.Error) as exc:
            self.data_status_var.set(f"Impossible de lire {filepath.name}: {exc}")
            return

        columns = self.data_tree["columns"]
        for row in recent_rows:
            self.data_tree.insert("", tk.END, values=tuple(row.get(column, "") for column in columns))
        shown = len(recent_rows)
        self.data_status_var.set(
            f"{filepath.name} — {shown} dernières lignes affichées sur {row_count}. "
            "Ouvrez le fichier pour accéder à l'historique complet."
        )

    def _clear_data_preview(self):
        for item in self.data_tree.get_children():
            self.data_tree.delete(item)

    def open_selected_data_file(self):
        selection = self.data_file_tree.selection()
        if not selection:
            return
        filepath = Path(selection[0]).resolve()
        if not webbrowser.open(filepath.as_uri()):
            messagebox.showerror("Données", f"Impossible d'ouvrir {filepath}")

    def _build_input_settings_tab(self, parent):
        mode_panel = ttk.Frame(parent, padding=12)
        mode_panel.pack(fill="x")
        ttk.Label(mode_panel, text="Acquisition").pack(side="left", padx=(0, 8))
        self.acquisition_mode_var = tk.StringVar(
            value="Simulation" if self.config_data["inputs"].get("mode") == "simulation" else "Capteurs"
        )
        ttk.Combobox(
            mode_panel,
            textvariable=self.acquisition_mode_var,
            values=("Capteurs", "Simulation"),
            state="readonly",
            width=16,
        ).pack(side="left")

        ports_panel = ttk.LabelFrame(parent, text="Ports des équipements", padding=12)
        ports_panel.pack(fill="x", padx=12, pady=(0, 12))
        headers = (
            "Équipement", "Transport", "IP source UDP", "Port UDP",
            "Port série (COM)", "Débit série",
        )
        for column, heading in enumerate(headers):
            ttk.Label(ports_panel, text=heading).grid(row=0, column=column, sticky="w", padx=5, pady=(0, 6))

        self.input_setting_vars = {}
        self.input_setting_widgets = {}
        equipment_labels = (("octans", "Octans"), ("current", "Courant"), ("wind", "Vent"))
        for row, (equipment, label) in enumerate(equipment_labels, start=1):
            config = self.config_data["inputs"].get(equipment, {})
            transport = config.get("transport", "udp")
            active_port = config.get("port", "")
            udp_port = config.get("udp_port", active_port if transport == "udp" else "")
            serial_port = config.get("serial_port", active_port if transport == "serial" else "")
            values = {
                "transport": tk.StringVar(value="Série" if transport == "serial" else "UDP"),
                "udp_host": tk.StringVar(value=str(config.get("host", ""))),
                "udp_port": tk.StringVar(value=str(udp_port)),
                "serial_port": tk.StringVar(value=str(serial_port)),
                "baudrate": tk.StringVar(value=str(config.get("baudrate", 115200))),
            }
            self.input_setting_vars[equipment] = values
            ttk.Label(ports_panel, text=label).grid(row=row, column=0, sticky="w", padx=5, pady=5)
            transport_combo = ttk.Combobox(
                ports_panel, textvariable=values["transport"], values=("UDP", "Série"),
                state="readonly", width=12,
            )
            transport_combo.grid(row=row, column=1, sticky="ew", padx=5, pady=5)
            host_entry = ttk.Entry(ports_panel, textvariable=values["udp_host"], width=18)
            host_entry.grid(row=row, column=2, sticky="ew", padx=5, pady=5)
            udp_entry = ttk.Entry(ports_panel, textvariable=values["udp_port"], width=12)
            udp_entry.grid(row=row, column=3, sticky="ew", padx=5, pady=5)
            serial_entry = ttk.Entry(ports_panel, textvariable=values["serial_port"], width=16)
            serial_entry.grid(row=row, column=4, sticky="ew", padx=5, pady=5)
            baudrate_entry = ttk.Entry(ports_panel, textvariable=values["baudrate"], width=12)
            baudrate_entry.grid(row=row, column=5, sticky="ew", padx=5, pady=5)
            self.input_setting_widgets[equipment] = {
                "udp_host": host_entry,
                "udp_port": udp_entry,
                "serial_port": serial_entry,
                "baudrate": baudrate_entry,
            }
            transport_combo.bind(
                "<<ComboboxSelected>>",
                lambda event, key=equipment: self._update_port_field_states(key),
            )
            self._update_port_field_states(equipment)

        for column in range(len(headers)):
            ports_panel.columnconfigure(column, weight=1 if column > 0 else 0)
        ttk.Label(
            parent,
            text="L'IP source UDP est facultative (vide = toute adresse). Les champs du transport inactif sont conservés.",
        ).pack(anchor="w", padx=16)
        actions = ttk.Frame(parent, padding=12)
        actions.pack(fill="x")
        ttk.Button(
            actions, text="Enregistrer et appliquer", command=self.save_input_settings
        ).pack(side="left")
        self.input_settings_status_var = tk.StringVar(value="")
        ttk.Label(actions, textvariable=self.input_settings_status_var).pack(side="left", padx=12)

    def _update_port_field_states(self, equipment):
        is_udp = self.input_setting_vars[equipment]["transport"].get() == "UDP"
        widgets = self.input_setting_widgets[equipment]
        widgets["udp_host"].configure(state="normal" if is_udp else "disabled")
        widgets["udp_port"].configure(state="normal" if is_udp else "disabled")
        widgets["serial_port"].configure(state="disabled" if is_udp else "normal")
        widgets["baudrate"].configure(state="disabled" if is_udp else "normal")

    def save_input_settings(self):
        equipment_settings = {}
        for equipment, values in self.input_setting_vars.items():
            equipment_settings[equipment] = {
                "transport": "serial" if values["transport"].get() == "Série" else "udp",
                "udp_host": values["udp_host"].get(),
                "udp_port": values["udp_port"].get(),
                "serial_port": values["serial_port"].get(),
                "baudrate": values["baudrate"].get(),
            }

        try:
            updated_config = copy.deepcopy(self.config_data)
            updated_config["inputs"] = build_input_configuration(
                self.config_data["inputs"],
                self.acquisition_mode_var.get(),
                equipment_settings,
            )
        except ValueError as exc:
            messagebox.showerror("Configuration des ports", str(exc))
            return

        config_path = BASE / "config.json"
        temporary_path = config_path.with_name("config.json.tmp")
        try:
            temporary_path.write_text(
                json.dumps(updated_config, indent=2, ensure_ascii=False) + "\n",
                encoding="utf-8",
            )
            temporary_path.replace(config_path)
        except OSError as exc:
            temporary_path.unlink(missing_ok=True)
            messagebox.showerror("Configuration des ports", f"Impossible d'enregistrer config.json: {exc}")
            return

        self.stop_inputs()
        self.config_data = updated_config
        with self.data_lock:
            self.heave.clear()
        self.start_inputs()
        self.input_settings_status_var.set("Configuration enregistrée et acquisition relancée.")
        self._update_rao_status()

    def _load_vessel_rao(self, vessel_name):
        profile = self.vessels[vessel_name]
        filepath = profile.get("rao_file")
        if filepath and not Path(filepath).is_absolute():
            filepath = BASE / filepath
        return RAOManager(filepath, profile.get("default_heading_deg", 0))

    def _configure_heading_control(self):
        headings = self.rao.available_headings
        values = [f"{heading:g}°" for heading in headings]
        self.heading_combo.configure(values=values, state="readonly" if values else "disabled")

        if headings:
            requested = self.vessels[self.selected_vessel].get("default_heading_deg", headings[0])
            heading = min(headings, key=lambda value: abs(value - float(requested)))
            self.rao.set_heading(heading)
            self.heading_var.set(f"{heading:g}°")
            self._update_rao_status()
        elif self.rao.load_error:
            self.heading_var.set("Indisponible")
            self._update_rao_status()
        else:
            self.heading_var.set("Indisponible")
            self._update_rao_status()

    def _update_rao_status(self):
        if self.config_data["inputs"].get("mode") == "simulation":
            self.rao_status_var.set("RAO non appliquée en simulation")
        elif self.rao.load_error:
            self.rao_status_var.set(f"RAO indisponible: {self.rao.load_error}")
        elif self.rao.available_headings:
            self.rao_status_var.set(
                f"RAO chargée ({self.rao.heading_deg:g}°, {len(self.rao.available_headings)} directions)"
            )
        else:
            self.rao_status_var.set("Aucune correction RAO")

    def on_vessel_select(self, event=None):
        vessel_name = self.vessel_var.get()
        if vessel_name not in self.vessels:
            return
        self.selected_vessel = vessel_name
        self.rao = self._load_vessel_rao(vessel_name)
        if hasattr(self, "processor"):
            self.processor.rao_manager = self.rao
        self._configure_heading_control()

    def on_heading_select(self, event=None):
        value = self.heading_var.get().replace("°", "")
        try:
            self.rao.set_heading(float(value))
        except ValueError:
            return
        self._update_rao_status()

    def start_inputs(self):
        self.workers = []
        mode = self.config_data["inputs"].get("mode", "simulation")
        
        if mode == "simulation":
            preset = self.config_data["inputs"].get("preset", "moderate")
            sim = Simulator(self.queue, self.sample_rate, preset=preset)
            self.workers.append(sim)
            self.status_var.set(f"Simulation active ({preset})")
        else:
            # Exemple de lancement hybride UDP/Série basé sur la config
            for sensor, cfg in self.config_data["inputs"].items():
                if sensor in ("mode", "preset"): continue
                is_octans = sensor == "octans"
                parser = (ExailParser.phlin if is_octans else
                          NMEAParser.wind if sensor == "wind" else NMEAParser.current)
                if cfg["transport"] == "udp":
                    w = UDPReceiver(
                        cfg["port"],
                        self.queue,
                        parser,
                        cfg.get("local_host", "0.0.0.0"),
                        cfg.get("host"),
                    )
                elif cfg["transport"] == "serial":
                    stream_parser = PhlinStreamParser() if is_octans else None
                    w = SerialReceiver(cfg["port"], cfg["baudrate"], self.queue, parser, stream_parser)
                self.workers.append(w)
            self.status_var.set("Acquisition capteurs active")

        for w in self.workers:
            w.start()

        # Démarrage du thread de calcul lourd
        self.processor = WaveProcessor(self.data_lock, self.heave, self.sample_rate, self.result_queue, self.config_data, self.rao)
        self.processor.start()

    def stop_inputs(self):
        workers = getattr(self, "workers", [])
        for worker in workers:
            worker.stop()
        processor = getattr(self, "processor", None)
        if processor:
            processor.stop()
        for worker in workers:
            worker.join(timeout=1.0)
        if processor:
            processor.join(timeout=2.5)

    def tick(self):
        """Boucle principale UI, extrêmement légère."""
        error_count = 0
        missing_sensors = 0

        while True:
            try:
                item = self.queue.get_nowait()
                kind = item[0]
                if kind == "error":
                    error_count += 1
                    msg = str(item[1])
                    self.logger.write_error("ERROR", "sensor", msg)
                    self.error_log_messages.append(msg)
                    self._append_error_log(msg)
                    continue
                if kind == "heave":
                    with self.data_lock:
                        self.heave.append(float(item[1]))
                elif kind == "current":
                    self.current.update(speed=float(item[1]), direction=float(item[2]), depth=float(item[3]))
                elif kind == "wind":
                    self.wind.update(speed=float(item[1]), direction=float(item[2]))
            except queue.Empty:
                break

        if len(self.heave) == 0:
            missing_sensors += 1
        if not np.isfinite(self.current["speed"]):
            missing_sensors += 1
        if not np.isfinite(self.wind["speed"]):
            missing_sensors += 1

        try:
            while True:
                self.wave = self.result_queue.get_nowait()
                self.history_t.append(datetime.now().strftime("%H:%M:%S"))
                self.history_hs.append(self.wave.hs_m if self.wave.valid else np.nan)
                self.history_tp.append(self.wave.tp_s if self.wave.valid else np.nan)
                self.history_current.append(
                    self.current["speed"] if np.isfinite(self.current["speed"]) else np.nan
                )
                self.history_wind.append(
                    self.wind["speed"] if np.isfinite(self.wind["speed"]) else np.nan
                )
                self.update_dashboard(error_count=error_count, missing_sensors=missing_sensors)
                self._log_current_state()
        except queue.Empty:
            pass

        if not self.history_hs:
            self.update_dashboard(error_count=error_count, missing_sensors=missing_sensors)

        self._update_sensor_status_summary(error_count=error_count, missing_sensors=missing_sensors)
        self.after(self.config_data["app"]["refresh_ms"], self.tick)

    def _log_current_state(self):
        heave_val = self.heave[-1] if self.heave else None
        timestamp_utc = datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")
        self.logger.write((
            timestamp_utc,
            heave_val,
            self.current["speed"], self.current["direction"], self.current["depth"],
            self.wind["speed"], self.wind["direction"],
            self.wave.hs_m if self.wave.valid else None,
            self.wave.tp_s if self.wave.valid else None
        ))

    def on_sensor_select(self, event=None):
        selection = self.sensor_listbox.curselection()
        if not selection:
            return
        name = self.sensor_listbox.get(selection[0])
        self.sensor_details.set(f"{name}: {self.sensor_status_vars.get(name, 'UNKNOWN')}")

    def _append_error_log(self, message):
        self.error_log_box.configure(state="normal")
        self.error_log_box.insert(tk.END, message + "\n")
        self.error_log_box.see(tk.END)
        self.error_log_box.configure(state="disabled")

    def _update_sensor_status_summary(self, error_count=0, missing_sensors=0):
        sensor_status = {
            "Octans": "GOOD" if len(self.heave) > 0 else "POOR",
            "Courant": "GOOD" if np.isfinite(self.current["speed"]) else "POOR",
            "Vent": "GOOD" if np.isfinite(self.wind["speed"]) else "POOR",
        }

        if error_count > 0:
            sensor_status["Octans"] = "MODERATE" if sensor_status["Octans"] != "POOR" else "POOR"

        for name, status in sensor_status.items():
            self.sensor_status_vars[name] = status
            self.sensor_listbox.itemconfig(self.sensor_listbox.get(0, tk.END).index(name), bg={"GOOD": "#c8e6c9", "MODERATE": "#ffe0b2", "POOR": "#ffcdd2"}[status])

        selected = self.sensor_listbox.curselection()
        if selected:
            name = self.sensor_listbox.get(selected[0])
            self.sensor_details.set(f"{name}: {sensor_status.get(name, 'UNKNOWN')}")

    def update_dashboard(self, error_count=0, missing_sensors=0):
        self.vars["Hs"][0].set(f"{self.wave.hs_m:.2f}" if self.wave.valid else "—")
        self.vars["Tp"][0].set(f"{self.wave.tp_s:.1f}" if self.wave.valid else "—")
        self.vars["Courant"][0].set(
            f"{self.current['speed']:.2f}" if np.isfinite(self.current["speed"]) else "—"
        )
        self.vars["Vent"][0].set(
            f"{self.wind['speed']:.1f}" if np.isfinite(self.wind["speed"]) else "—"
        )

        quality = determine_reception_quality(self.wave.valid, missing_sensors=missing_sensors, error_count=error_count)
        self.quality_var.set(f"{quality} ({missing_sensors} missing / {error_count} errors)")
        colors = {"GOOD": ("#2e7d32", "white"), "MODERATE": ("#f9a825", "black"), "POOR": ("#c62828", "white")}
        bg, fg = colors.get(quality, ("#2e7d32", "white"))
        self.quality_label.configure(background=bg, foreground=fg, padding=(12, 4))

        a = self.config_data["alerts"]
        alarms = []
        if self.wave.valid and self.wave.hs_m > a["hs_max_m"]: alarms.append("Hs")
        if np.isfinite(self.wind["speed"]) and self.wind["speed"] > a["wind_max_kn"]: alarms.append("Vent")
        if quality == "POOR": alarms.append("Reception")
        if quality == "MODERATE": alarms.append("Qualité faible")

        self.alert_var.set("ALERTE: " + ", ".join(alarms) if alarms else "NORMAL")

        self.wave_ax.clear()
        self.tp_ax.clear()
        self.environment_ax.clear()
        self.wind_ax.clear()
        x = list(self.history_t)
        hs_line, = self.wave_ax.plot(x, self.history_hs, color="#2878b5", label="Hs (m)")
        tp_line, = self.tp_ax.plot(x, self.history_tp, color="#d95f02", label="Tp (s)")
        self.wave_ax.set_ylabel("Hs (m)", color="#2878b5")
        self.tp_ax.set_ylabel("Tp (s)", color="#d95f02")
        self.wave_ax.set_title("État de mer")
        self.wave_ax.grid(True, alpha=0.25)
        self.wave_ax.legend((hs_line, tp_line), ("Hs (m)", "Tp (s)"), loc="upper left")

        current_line, = self.environment_ax.plot(
            x, self.history_current, color="#238b45", label="Courant (kn)"
        )
        wind_line, = self.wind_ax.plot(x, self.history_wind, color="#b23a48", label="Vent (kn)")
        self.environment_ax.set_ylabel("Courant (kn)", color="#238b45")
        self.wind_ax.set_ylabel("Vent (kn)", color="#b23a48")
        self.environment_ax.set_title("Vent et courant")
        self.environment_ax.grid(True, alpha=0.25)
        self.environment_ax.legend(
            (current_line, wind_line), ("Courant (kn)", "Vent (kn)"), loc="upper left"
        )
        if len(x) > 1:
            step = max(1, len(x) // 8)
            self.environment_ax.set_xticks(range(0, len(x), step))
            self.environment_ax.set_xticklabels(x[::step], rotation=30)
        self.canvas.draw_idle()

    def on_close(self):
        self.stop_inputs()
        self.logger.close()
        self.destroy()

if __name__ == "__main__":
    App().mainloop()