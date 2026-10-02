import csv
import queue
import socket
import struct
import threading

import numpy as np
import pytest
from openpyxl import Workbook
from app import (
    ExailParser,
    PhlinStreamParser,
    RAOManager,
    SEA_STATE_PRESETS,
    CSVLogger,
    StdBinStreamParser,
    UDPReceiver,
    WaveProcessor,
    build_input_configuration,
    compute_wave,
    determine_reception_quality,
    simulate_sea_state,
)


def make_phlin_frame(x, y, heave):
    body = f"PHLIN,{x:.3f},{y:.3f},{heave:.3f}"
    checksum = 0
    for character in body.encode("ascii"):
        checksum ^= character
    return f"${body}*{checksum:02X}\r\n".encode("ascii")


def make_stdbin_frame(heave):
    header = b"IX" + bytes([3])
    header += (1 << 2).to_bytes(4, "big")
    header += b"\x00" * 8
    header += (41).to_bytes(2, "big")
    header += b"\x00" * 8
    return header + struct.pack(">4f", 1.0, heave, 3.0, 4.0)


def test_stdbin_parser_extracts_heave_from_fragmented_stream():
    frame = make_stdbin_frame(2.5)
    parser = StdBinStreamParser()

    assert parser.feed(b"noise" + frame[:12]) == []
    assert parser.feed(frame[12:]) == [("heave", 2.5)]
    assert ExailParser.stdbin(frame) == ("heave", 2.5)


def test_phlin_parser_extracts_heave_from_fragmented_stream():
    frame = make_phlin_frame(1.234, 2.345, -3.456)
    parser = PhlinStreamParser()

    assert parser.feed(frame[:10]) == []
    assert parser.feed(frame[10:]) == [("heave", -3.456)]
    assert ExailParser.phlin(frame) == ("heave", -3.456)


def test_phlin_parser_extracts_heave_from_multisentence_udp_datagram():
    payload = (
        b"$HEHDT,84.44,T*23\r\n"
        b"$PHTRO,0.41,P,0.44,B*46\r\n"
        b"$PHLIN,0.009,-0.063,0.013*72\r\n"
        b"$PHSPD,0.003,-0.008,-0.008*5E\r\n"
        b"$PHCMP,3611.97,N,0.00,N*7D\r\n"
        b"$PHINF,03000000*76\r\n"
    )

    assert ExailParser.phlin(payload) == ("heave", 0.013)


def test_udp_receiver_delivers_heave_from_multisentence_datagram():
    payload = (
        b"$HEHDT,84.44,T*23\r\n"
        b"$PHTRO,0.41,P,0.44,B*46\r\n"
        b"$PHLIN,0.009,-0.063,0.013*72\r\n"
        b"$PHSPD,0.003,-0.008,-0.008*5E\r\n"
    )
    with socket.socket(socket.AF_INET, socket.SOCK_DGRAM) as reservation:
        reservation.bind(("127.0.0.1", 0))
        port = reservation.getsockname()[1]

    messages = queue.Queue()
    receiver = UDPReceiver(
        port,
        messages,
        ExailParser.phlin,
        bind_host="127.0.0.1",
        source_host="127.0.0.1",
    )
    receiver.start()
    try:
        with socket.socket(socket.AF_INET, socket.SOCK_DGRAM) as sender:
            received = None
            for _ in range(10):
                sender.sendto(payload, ("127.0.0.1", port))
                try:
                    received = messages.get(timeout=0.1)
                    break
                except queue.Empty:
                    continue
        assert received is not None
        assert received[:2] == ("heave", 0.013)
    finally:
        receiver.stop()
        receiver.join(timeout=1.0)


def test_phlin_parser_rejects_invalid_checksum():
    frame = make_phlin_frame(1.234, 2.345, 3.456)
    invalid_frame = frame.replace(b"*", b"*00", 1)

    assert ExailParser.phlin(invalid_frame) is None


def test_wave_returns_valid_result():
    fs = 5.0
    t = np.arange(0, 20*60, 1/fs)
    x = 0.5*np.sin(2*np.pi*0.1*t)
    r = compute_wave(x, fs)
    assert r.valid
    assert r.hs_m > 0
    assert abs(r.tp_s - 10) < 1.0


def test_simulated_wave_result_is_available_after_ten_seconds():
    sample_rate = 5.0
    time_s = np.arange(0, 10, 1 / sample_rate)
    heave = 0.5 * np.sin(2 * np.pi * 0.1 * time_s)

    result = compute_wave(heave, sample_rate, min_duration_seconds=10)

    assert result.valid
    assert abs(result.tp_s - 10) < 1.0


def test_wave_processor_skips_vessel_rao_for_simulated_data():
    class LowGainRAO:
        def get_gain(self, frequency):
            return 0.01

    sample_rate = 5.0
    time_s = np.arange(0, 10, 1 / sample_rate)
    heave = 0.5 * np.sin(2 * np.pi * 0.1 * time_s)
    results = queue.Queue()
    processor = WaveProcessor(
        threading.Lock(),
        heave,
        sample_rate,
        results,
        {"inputs": {"mode": "simulation"}, "wave": {"min_frequency_hz": 0.03, "max_frequency_hz": 0.5}},
        LowGainRAO(),
    )
    processor.start()
    try:
        result = results.get(timeout=2)
    finally:
        processor.stop()
        processor.join(timeout=1)

    assert result.valid
    assert result.hs_m < 1.0
    assert abs(result.tp_s - 10) < 1.0


def test_rao_manager_loads_heading_specific_excel_curves(tmp_path):
    workbook = Workbook()
    for heading, gains in (("0", (0.8, 0.4)), ("90", (0.6, 0.2))):
        sheet = workbook.active if heading == "0" else workbook.create_sheet(heading)
        sheet.title = heading
        sheet.cell(row=5, column=2, value=2.0)
        sheet.cell(row=5, column=5, value=gains[1])
        sheet.cell(row=6, column=2, value=4.0)
        sheet.cell(row=6, column=5, value=gains[0])
    filepath = tmp_path / "rao_iot.xlsx"
    workbook.save(filepath)

    rao = RAOManager(filepath)

    assert rao.available_headings == (0.0, 90.0)
    assert rao.get_gain(0.25) == 0.8
    assert rao.get_gain(0.1) == 1.0
    rao.set_heading(90)
    assert rao.get_gain(0.25) == 0.6


def test_rao_manager_loads_molene_excel_layout(tmp_path):
    workbook = Workbook()
    for heading, gain in (("0°", 0.8), ("90°", 0.6)):
        sheet = workbook.active if heading == "0°" else workbook.create_sheet(heading)
        sheet.title = heading
        sheet.cell(row=4, column=1, value="Period")
        sheet.cell(row=4, column=6, value="Heave")
        sheet.cell(row=7, column=1, value=4.0)
        sheet.cell(row=7, column=6, value=gain)
    filepath = tmp_path / "rao_molene.xlsx"
    workbook.save(filepath)

    rao = RAOManager(filepath)

    assert rao.available_headings == (0.0, 90.0)
    assert rao.get_gain(0.25) == 0.8
    rao.set_heading(90)
    assert rao.get_gain(0.25) == 0.6


def test_simulate_sea_state_has_realistic_variability():
    samples = [simulate_sea_state(i * 0.2) for i in range(600)]
    arr = np.asarray(samples, dtype=float)
    assert arr.std() > 0.1
    assert np.ptp(arr) > 0.5


def test_sea_state_presets_are_defined():
    assert set(SEA_STATE_PRESETS) == {"calm", "moderate", "rough", "instrument"}
    for key, cfg in SEA_STATE_PRESETS.items():
        assert cfg["swell_h"] > 0
        assert cfg["noise_std"] >= 0


def test_csv_log_contains_measurements_and_errors(tmp_path):
    logger = CSVLogger(tmp_path)
    logger.start()
    logger.write(("2026-10-01T12:00:00Z", 0.4, 1.2, 90.0, 20.0, 10.0, 180.0, 1.5, 8.0))
    logger.write_error("ERROR", "sensor", "Test, message")
    logger.close()

    with logger.log_path.open(encoding="utf-8", newline="") as source:
        rows = list(csv.DictReader(source))

    assert len(rows) == 2
    assert logger.log_path.suffix == ".csv"
    assert rows[0]["type"] == "measurement"
    assert rows[0]["heave_m"] == "0.4"
    assert rows[1]["type"] == "error"
    assert rows[1]["message"] == "Test, message"


def test_build_input_configuration_preserves_udp_and_serial_ports():
    current = {
        "mode": "udp",
        "octans": {"transport": "udp", "port": 9998, "local_host": "0.0.0.0"},
    }
    updated = build_input_configuration(
        current,
        "Capteurs",
        {
            "octans": {
                "transport": "serial",
                "udp_port": "9998",
                "serial_port": "COM4",
                "baudrate": "115200",
            }
        },
    )

    assert current["octans"]["transport"] == "udp"
    assert updated["mode"] == "udp"
    assert updated["octans"]["port"] == "COM4"
    assert updated["octans"]["udp_port"] == 9998
    assert updated["octans"]["serial_port"] == "COM4"
    assert updated["octans"]["baudrate"] == 115200


def test_build_input_configuration_persists_optional_udp_source_ip():
    updated = build_input_configuration(
        {"mode": "udp", "wind": {"transport": "udp", "port": 5002}},
        "Capteurs",
        {
            "wind": {
                "transport": "udp",
                "udp_host": "192.168.10.45",
                "udp_port": "5002",
                "serial_port": "",
                "baudrate": "115200",
            }
        },
    )

    assert updated["wind"]["host"] == "192.168.10.45"
    assert updated["wind"]["port"] == 5002


def test_build_input_configuration_rejects_invalid_udp_source_ip():
    with pytest.raises(ValueError, match="IPv4 valide"):
        build_input_configuration(
            {"mode": "udp", "wind": {"transport": "udp", "port": 5002}},
            "Capteurs",
            {
                "wind": {
                    "transport": "udp",
                    "udp_host": "not-an-ip",
                    "udp_port": "5002",
                    "serial_port": "",
                    "baudrate": "115200",
                }
            },
        )


def test_build_input_configuration_rejects_invalid_ports():
    with pytest.raises(ValueError, match="1 et 65535"):
        build_input_configuration(
            {"mode": "udp", "wind": {"transport": "udp", "port": 5002}},
            "Capteurs",
            {"wind": {"transport": "udp", "udp_port": "70000", "serial_port": "", "baudrate": "115200"}},
        )


def test_reception_quality_levels_are_defined():
    assert determine_reception_quality(valid=True, missing_sensors=0, error_count=0) == "GOOD"
    assert determine_reception_quality(valid=False, missing_sensors=2, error_count=3) == "POOR"
    assert determine_reception_quality(valid=True, missing_sensors=1, error_count=1) == "MODERATE"
